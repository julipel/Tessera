"""Чтение табличных файлов источника `table` (CSV/XLSX) в строки значений ячеек."""

import csv
import zipfile
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO, StringIO

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from app.modules.knowledge.infrastructure.file_parsers import (
    DEFAULT_MAX_UNPACKED_BYTES,
    decode_text,
)

# CSV даёт только строки; XLSX — типизированные значения (формулы — последним значением).
type Cell = str | int | float | bool | datetime | date | time | None
type Row = Sequence[Cell]
type TableReader = Callable[[bytes, "ReadOptions"], list[Row]]

_CSV_DELIMITERS = ",;\t|"


class TableReadError(Exception):
    """Файл повреждён, зашифрован или в нём нет нужного листа."""


@dataclass(frozen=True, slots=True)
class ReadOptions:
    sheet: str | None = None
    delimiter: str | None = None
    max_unpacked_bytes: int = DEFAULT_MAX_UNPACKED_BYTES


def read_csv(data: bytes, options: ReadOptions) -> list[Row]:
    """Кодировка — как у текстовых файлов; разделитель из конфига или угадывается."""
    text = decode_text(data)
    delimiter = options.delimiter or _sniff_delimiter(text)
    try:
        return list(csv.reader(StringIO(text), delimiter=delimiter))
    except csv.Error as e:
        raise TableReadError(f"не удалось прочитать CSV: {e}") from e


def _sniff_delimiter(text: str) -> str:
    """Самый частый разделитель (`,` `;` `|`, табуляция) в строке заголовка. `csv.Sniffer`
    ошибается на заголовках вида «Объём, мл» и на файлах из одной строки. Ничья — запятая."""
    header = next((line for line in text.split("\n") if line.strip()), "")
    counts = {d: header.count(d) for d in _CSV_DELIMITERS}
    best = max(counts, key=lambda d: counts[d])
    return best if counts[best] > counts[","] else ","


def read_xlsx(data: bytes, options: ReadOptions) -> list[Row]:
    """Лист из конфига или первый; ячейки с формулами — сохранённым значением."""
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            unpacked = sum(info.file_size for info in archive.infolist())
        if unpacked > options.max_unpacked_bytes:
            raise TableReadError(
                f"XLSX распаковывается больше чем в {options.max_unpacked_bytes} байт"
            )
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            if options.sheet is None:
                sheet = workbook.worksheets[0]
            elif options.sheet in workbook.sheetnames:
                sheet = workbook[options.sheet]
            else:
                raise TableReadError(f"в XLSX нет листа «{options.sheet}»")
            if not isinstance(sheet, ReadOnlyWorksheet):
                raise TableReadError("лист XLSX не содержит таблицы")
            return [tuple(map(_xlsx_cell, row)) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    # Битый XML внутри архива — SyntaxError (в т.ч. от lxml), прочее — ошибки zip/openpyxl.
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, SyntaxError) as e:
        raise TableReadError(f"не удалось прочитать XLSX: {e}") from e


def _xlsx_cell(value: object) -> Cell:
    """Значения openpyxl, которых нет в `Cell` (Decimal, timedelta…), — строкой."""
    if value is None or isinstance(value, str | int | float | datetime | date | time):
        return value
    return str(value)


READERS: dict[str, TableReader] = {".csv": read_csv, ".xlsx": read_xlsx}
