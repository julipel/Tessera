"""Парсеры PDF/DOCX коннектора `file`: приведение к markdown, заголовки, ошибки."""

import zipfile
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import docx
import pytest
from pypdf import PdfReader, PdfWriter

from app.modules.knowledge.infrastructure.file_parsers import (
    FileParseError,
    parse_docx,
    parse_pdf,
)
from app.modules.knowledge.public import (
    DocumentItem,
    FileConnector,
    MarkdownChunker,
    SourceFileError,
    SourceSpec,
)
from app.modules.shared.public import TenantId


def make_pdf(pages: list[list[str]], title: str | None = None) -> bytes:
    """Минимальный PDF: строки страницы — Helvetica 12pt сверху вниз (только ASCII)."""
    n = len(pages)
    page_ids = [4 + 2 * i for i in range(n)]
    objects: dict[int, bytes] = {
        1: b"<< /Type /Catalog /Pages 2 0 R >>",
        2: b"<< /Type /Pages /Kids [%s] /Count %d >>"
        % (b" ".join(b"%d 0 R" % i for i in page_ids), n),
        3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    }
    for page_id, lines in zip(page_ids, pages, strict=True):
        ops = [b"BT /F1 12 Tf 72 720 Td 14 TL"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(b"(%s) Tj T*" % escaped.encode("ascii"))
        ops.append(b"ET")
        stream = b"\n".join(ops)
        objects[page_id] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % (page_id + 1)
        )
        objects[page_id + 1] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)
    info_id = None
    if title is not None:
        info_id = max(objects) + 1
        objects[info_id] = b"<< /Title (%s) >>" % title.encode("ascii")

    out = BytesIO(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for obj_id in sorted(objects):
        offsets[obj_id] = out.tell()
        out.write(b"%d 0 obj\n%s\nendobj\n" % (obj_id, objects[obj_id]))
    xref = out.tell()
    size = max(objects) + 1
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % size)
    for obj_id in range(1, size):
        out.write(b"%010d 00000 n \n" % offsets[obj_id])
    info = b" /Info %d 0 R" % info_id if info_id else b""
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R%s >>\nstartxref\n%d\n%%%%EOF\n" % (size, info, xref)
    )
    return out.getvalue()


def encrypt(pdf: bytes, user_password: str) -> bytes:
    writer = PdfWriter(clone_from=PdfReader(BytesIO(pdf)))
    writer.encrypt(user_password, owner_password="owner", algorithm="RC4-128")
    out = BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx(title: str | None = None) -> bytes:
    document = docx.Document()
    if title is not None:
        document.core_properties.title = title
    document.add_heading("Условия", level=0)
    document.add_paragraph("Вводный абзац.")
    document.add_heading("Доставка", level=1)
    document.add_paragraph("Курьером по городу.")
    document.add_heading("Сроки", level=2)
    document.add_paragraph("# не заголовок, а текст")
    document.add_paragraph("Один день", style="List Bullet")
    document.add_paragraph("Два дня", style="List Number")
    table = document.add_table(rows=2, cols=2)
    for (r, c), value in {(0, 0): "Зона", (0, 1): "Цена", (1, 0): "A|B", (1, 1): "300"}.items():
        table.cell(r, c).text = value
    document.add_paragraph("")
    out = BytesIO()
    document.save(out)
    return out.getvalue()


# --- PDF ---


def test_pdf_pages_become_paragraphs_with_metadata_title() -> None:
    pdf = make_pdf(
        [["Delivery terms", "Couriers work daily."], ["Returns within 14 days."]], "Terms"
    )
    parsed = parse_pdf(pdf)
    assert parsed.title == "Terms"
    assert parsed.text == "Delivery terms\nCouriers work daily.\n\nReturns within 14 days."


def test_pdf_joins_hyphenated_words_and_escapes_markup() -> None:
    parsed = parse_pdf(
        make_pdf([["Detailed infor-", "mation here.", "# 1 in sales", "Self-", "Made"]])
    )
    assert parsed.title is None
    assert parsed.text == "Detailed information here.\n\\# 1 in sales\nSelf-\nMade"
    chunks = MarkdownChunker().chunk(parsed.text)
    assert [c.section for c in chunks] == [None]


def test_pdf_without_text_layer_is_an_error() -> None:
    with pytest.raises(FileParseError, match="текстового слоя"):
        parse_pdf(make_pdf([[]]))


def test_pdf_with_password_is_an_error_but_empty_user_password_is_opened() -> None:
    pdf = make_pdf([["Secret price list"]])
    with pytest.raises(FileParseError, match="зашифрован"):
        parse_pdf(encrypt(pdf, "secret"))
    assert parse_pdf(encrypt(pdf, "")).text == "Secret price list"


def test_garbage_pdf_is_an_error() -> None:
    with pytest.raises(FileParseError, match="PDF"):
        parse_pdf(b"not a pdf at all")


# --- DOCX ---


def test_docx_to_markdown() -> None:
    parsed = parse_docx(make_docx())
    assert parsed.title == "Условия"
    assert parsed.text == (
        "# Условия\n\n"
        "Вводный абзац.\n\n"
        "# Доставка\n\n"
        "Курьером по городу.\n\n"
        "## Сроки\n\n"
        "\\# не заголовок, а текст\n\n"
        "- Один день\n\n"
        "- Два дня\n\n"
        "| Зона | Цена |\n| --- | --- |\n| A\\|B | 300 |"
    )


def test_docx_headings_become_chunk_sections() -> None:
    chunks = MarkdownChunker().chunk(parse_docx(make_docx()).text)
    assert [c.section for c in chunks] == ["Условия", "Доставка", "Доставка > Сроки"]
    assert "| A\\|B | 300 |" in chunks[-1].text


def test_docx_core_properties_title_wins() -> None:
    assert parse_docx(make_docx(title="Оферта")).title == "Оферта"


def test_broken_docx_is_an_error() -> None:
    with pytest.raises(FileParseError, match="DOCX"):
        parse_docx(b"PK\x03\x04 broken")
    archive = BytesIO()
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("hello.txt", "не word")
    with pytest.raises(FileParseError, match="DOCX"):
        parse_docx(archive.getvalue())


def test_docx_unpacked_size_is_limited() -> None:
    with pytest.raises(FileParseError, match="распаковывается"):
        parse_docx(make_docx(), max_unpacked_bytes=1000)


# --- коннектор ---


async def test_connector_reads_pdf_and_docx(tmp_path: Path) -> None:
    source = SourceSpec(TenantId(uuid4()), uuid4(), {})
    connector = FileConnector(tmp_path)
    base = connector.source_dir(source)
    base.mkdir(parents=True)
    (base / "price.pdf").write_bytes(make_pdf([["Price list"]]))
    (base / "offer.DOCX").write_bytes(make_docx())
    (base / "legacy.doc").write_bytes(b"old word")
    (base / "broken.pdf").write_bytes(b"garbage")

    listing = await connector.discover(source)
    assert [r.external_id for r in listing.refs] == ["broken.pdf", "offer.DOCX", "price.pdf"]

    pdf = await connector.fetch(source, listing.refs[2])
    assert isinstance(pdf, DocumentItem)
    assert (pdf.title, pdf.text, pdf.metadata["format"]) == ("price", "Price list", "pdf")
    offer = await connector.fetch(source, listing.refs[1])
    assert isinstance(offer, DocumentItem)
    assert offer.title == "Условия"
    assert offer.metadata["format"] == "docx"
    with pytest.raises(SourceFileError, match=r"broken\.pdf: не удалось прочитать PDF"):
        await connector.fetch(source, listing.refs[0])
